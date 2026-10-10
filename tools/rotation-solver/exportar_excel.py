"""Exportação pra Excel (openpyxl) — SPEC.md, secção 7.

Uso:
    python exportar_excel.py saida.xlsx config1.json:cenario1.json [config2.json:cenario2.json ...]

Cada par é "config.json:plano_resolvido.json" (o plano já gravado por
gerar.py em saida/<nome>.json). O nome do cenário na planilha vem do nome do
ficheiro de config (sem extensão).
"""
from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from relatorio import construir_quartetos, estatisticas_atleta

COR_CABECALHO = "18364C"
COR_TEXTO = "183042"
COR_LINHA_ALTERNADA = "EEF2F5"
COR_GRADE_A = "DDEAF4"
COR_GRADE_B = "E8EFEB"
COR_BORDA_GROSSA = "4A4A4A"
COR_BORDA_FINA = "B8C4CE"

FONTE_CABECALHO = Font(color="FFFFFF", bold=True)
FONTE_TEXTO = Font(color=COR_TEXTO)
PREENCH_CABECALHO = PatternFill("solid", fgColor=COR_CABECALHO)
PREENCH_ALTERNADA = PatternFill("solid", fgColor=COR_LINHA_ALTERNADA)
FORMATO_MIN = '0" min"'

BORDA_FINA = Side(style="thin", color=COR_BORDA_FINA)
BORDA_GROSSA = Side(style="medium", color=COR_BORDA_GROSSA)


def _ajustar_pagina(ws: Worksheet, paisagem: bool = False) -> None:
    """Sem isto, uma folha com colunas largadas (ex.: `column_dimensions` pra
    mais colunas do que as que têm conteúdo de verdade) imprime/exporta pra
    PDF com páginas extra em branco à direita — `fitToWidth=1` força tudo a
    caber na largura real do conteúdo, nunca sobra página vazia."""
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if paisagem:
        ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _cabecalho(ws: Worksheet, linha: int, col_ini: int, titulos: list[str]) -> None:
    # wrap_text pra títulos longos (ex.: "<cenário> — máx. seguido") nunca cortarem
    # visualmente — melhor quebrar em 2 linhas que truncar sem aviso nenhum.
    for i, titulo in enumerate(titulos):
        c = ws.cell(row=linha, column=col_ini + i, value=titulo)
        c.font = FONTE_CABECALHO
        c.fill = PREENCH_CABECALHO
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")
    if ws.row_dimensions[linha].height is None or ws.row_dimensions[linha].height < 30:
        ws.row_dimensions[linha].height = 30


def _linha_zebra(ws: Worksheet, linha: int, col_ini: int, col_fim: int) -> None:
    if linha % 2 == 0:
        for col in range(col_ini, col_fim + 1):
            ws.cell(row=linha, column=col).fill = PREENCH_ALTERNADA


# --------------------------------------------------------------------------- Registos


def escrever_registos(ws: Worksheet, cenarios: dict[str, tuple[dict, dict]]) -> dict:
    """Devolve linha_de[(cenario, parte, vaga, inicio_parte)] = numero da linha."""
    ws.title = "Registos"
    _ajustar_pagina(ws)
    ws.cell(row=1, column=1, value="Registos").font = Font(bold=True, size=14, color=COR_TEXTO)
    titulos = ["Cenário", "Parte", "Posição", "Atleta", "Início parte", "Fim parte", "Duração", "Início jogo", "Fim jogo"]
    _cabecalho(ws, 3, 1, titulos)
    ws.freeze_panes = "A4"

    linha_de: dict = {}
    r = 4
    for nome_cenario, (cfg, plano) in cenarios.items():
        minutos_por_parte = cfg["jogo"]["minutos_por_parte"]
        for parte_str in sorted(plano.keys(), key=int):
            parte = int(parte_str)
            offset = (parte - 1) * minutos_por_parte
            for vaga in cfg["posicoes"]:
                for nome, a, b in plano[parte_str].get(vaga, []):
                    ws.cell(row=r, column=1, value=nome_cenario)
                    ws.cell(row=r, column=2, value=parte)
                    ws.cell(row=r, column=3, value=vaga)
                    ws.cell(row=r, column=4, value=nome)
                    ws.cell(row=r, column=5, value=a)
                    ws.cell(row=r, column=6, value=b)
                    dur = ws.cell(row=r, column=7, value=f"=F{r}-E{r}")
                    dur.number_format = FORMATO_MIN
                    ws.cell(row=r, column=8, value=a + offset)
                    ws.cell(row=r, column=9, value=b + offset)
                    linha_de[(nome_cenario, parte, vaga, a)] = r
                    r += 1
    for col, largura in zip("ABCDEFGHI", [12, 7, 14, 14, 11, 10, 9, 11, 10]):
        ws.column_dimensions[col].width = largura
    ws.print_title_rows = "5:5"
    return linha_de


def _linha_ativa_em(linha_de: dict, cenario: str, parte: int, vaga: str, plano_parte: dict, minuto: int) -> int:
    for nome, a, b in plano_parte.get(vaga, []):
        if a <= minuto < b:
            return linha_de[(cenario, parte, vaga, a)]
    raise KeyError(f"Nenhum segmento de {vaga} cobre o minuto {minuto} ({cenario}, parte {parte}).")


# --------------------------------------------------------------------------- Cenário k


def escrever_cenario(ws: Worksheet, nome_cenario: str, descricao: str, cfg: dict, plano: dict, linha_de: dict) -> None:
    _ajustar_pagina(ws, paisagem=True)

    titulo = f"{nome_cenario}: {descricao}" if descricao and descricao != nome_cenario else nome_cenario
    ws.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=14, color=COR_TEXTO)
    ws.cell(row=2, column=1, value=_texto_legenda(cfg)).font = Font(italic=True, color=COR_TEXTO, size=9)

    linha_atual = 4
    for parte_str in sorted(plano.keys(), key=int):
        parte = int(parte_str)
        linha_atual = _escrever_grelha_parte(ws, linha_atual, nome_cenario, parte, cfg, plano[parte_str], linha_de)
        linha_atual += 1

    _escrever_tabela_quartetos(ws, linha_atual, nome_cenario, cfg, plano, linha_de)


def _texto_legenda(cfg: dict) -> str:
    r = cfg["regras"]
    return (
        f"Máx. seguido: {r['max_minutos_seguidos']} min · Entrada mín.: {r['min_minutos_por_entrada']} min · "
        f"Pausa mín.: {r['min_minutos_pausa']} min · Apoio em campo: máx. {r['max_apoio_em_campo']} · "
        f"Fecho sem apoio: últimos {r['fecho_sem_apoio_min']} min · Fecho com titulares: últimos {r['fecho_com_titulares_min']} min"
    )


def _escrever_grelha_parte(ws: Worksheet, linha_ini: int, nome_cenario: str, parte: int, cfg: dict, plan_parte: dict, linha_de: dict) -> int:
    N = cfg["jogo"]["minutos_por_parte"]
    COL_LABEL = 1
    COL_MIN0 = 2

    ws.cell(row=linha_ini, column=COL_LABEL, value=f"{parte}.ª parte — Minutos").font = Font(bold=True, color=COR_TEXTO)
    linha_minutos = linha_ini + 1
    ws.cell(row=linha_minutos, column=COL_LABEL, value="Minutos")
    for t in range(N + 1):
        c = ws.cell(row=linha_minutos, column=COL_MIN0 + t, value=t)
        c.alignment = Alignment(horizontal="center")
        c.font = Font(size=8, color=COR_TEXTO)
        ws.column_dimensions[get_column_letter(COL_MIN0 + t)].width = 5.6

    ordem_vagas = cfg["posicoes"]
    for i, vaga in enumerate(ordem_vagas):
        linha = linha_minutos + 1 + i
        ws.row_dimensions[linha].height = 48
        ws.cell(row=linha, column=COL_LABEL, value=vaga).font = Font(bold=True, color=COR_TEXTO)
        cor = COR_GRADE_A if i % 2 == 0 else COR_GRADE_B
        preench = PatternFill("solid", fgColor=cor)

        segs = plan_parte.get(vaga, [])
        for nome, a, b in segs:
            r_reg = linha_de[(nome_cenario, parte, vaga, a)]
            for t in range(a, b):
                col = COL_MIN0 + t
                c = ws.cell(row=linha, column=col)
                c.fill = preench
                c.alignment = Alignment(horizontal="left", wrap_text=False)
                esquerda = BORDA_GROSSA if t == a else BORDA_FINA
                direita = BORDA_GROSSA if t == b - 1 else BORDA_FINA
                c.border = Border(top=BORDA_GROSSA, bottom=BORDA_GROSSA, left=esquerda, right=direita)
                if t == a:
                    if (b - a) <= 2:
                        c.value = f"=Registos!D{r_reg}"
                    else:
                        c.value = f'=" "&Registos!D{r_reg}&" "&Registos!E{r_reg}&"–"&Registos!F{r_reg}'
                    c.font = FONTE_TEXTO

    return linha_minutos + 1 + len(ordem_vagas)


def _escrever_tabela_quartetos(ws: Worksheet, linha_ini: int, nome_cenario: str, cfg: dict, plano: dict, linha_de: dict) -> int:
    ws.cell(row=linha_ini, column=1, value="Quartetos").font = Font(bold=True, color=COR_TEXTO)
    linha_cab = linha_ini + 1
    ordem_vagas = cfg["posicoes"]
    titulos = ["Parte", "Minutos (parte)", "Minutos (jogo)"] + ordem_vagas + ["Entram", "Saem", "Mudam de posição"]
    _cabecalho(ws, linha_cab, 1, titulos)

    N = cfg["jogo"]["minutos_por_parte"]
    r = linha_cab + 1
    for parte_str in sorted(plano.keys(), key=int):
        parte = int(parte_str)
        offset = (parte - 1) * N
        linhas = construir_quartetos(plano[parte_str], N, ordem_vagas)
        for q in linhas:
            ws.cell(row=r, column=1, value=parte)
            rotulo_parte = "Quarteto inicial" if q["inicial"] else f"{q['inicio']}–{q['fim']}"
            ws.cell(row=r, column=2, value=rotulo_parte)
            ws.cell(row=r, column=3, value=f"{q['inicio']+offset}–{q['fim']+offset}")
            for i, vaga in enumerate(ordem_vagas):
                col = 4 + i
                try:
                    r_reg = _linha_ativa_em(linha_de, nome_cenario, parte, vaga, plano[parte_str], q["inicio"])
                    ws.cell(row=r, column=col, value=f"=Registos!D{r_reg}")
                except KeyError:
                    ws.cell(row=r, column=col, value="")
            ws.cell(row=r, column=8, value=", ".join(q["entram"]))
            ws.cell(row=r, column=9, value=", ".join(q["saem"]))
            ws.cell(row=r, column=10, value="; ".join(q["mudam_de_posicao"]))
            _linha_zebra(ws, r, 1, 10)
            r += 1
    for col, largura in zip("ABCDEFGHIJ", [7, 15, 15, 13, 13, 13, 13, 18, 18, 28]):
        ws.column_dimensions[col].width = largura
    return r


# --------------------------------------------------------------------------- Resumo


def escrever_resumo(ws: Worksheet, cenarios: dict[str, tuple[dict, dict]]) -> None:
    ws.title = "Resumo"
    _ajustar_pagina(ws, paisagem=True)
    ws.cell(row=1, column=1, value="Resumo").font = Font(bold=True, size=14, color=COR_TEXTO)

    nomes_cenarios = list(cenarios.keys())
    todos_atletas = sorted({a["nome"] for cfg, _ in cenarios.values() for a in cfg["atletas"]})

    # (1) minutos em campo por cenário.
    linha = 3
    ws.cell(row=linha, column=1, value="Minutos em campo por atleta").font = Font(bold=True, color=COR_TEXTO)
    linha += 1
    _cabecalho(ws, linha, 1, ["Atleta"] + nomes_cenarios)
    linha_cab1 = linha
    linha += 1
    for nome in todos_atletas:
        ws.cell(row=linha, column=1, value=nome)
        for i, nc in enumerate(nomes_cenarios):
            col = 2 + i
            col_letra = get_column_letter(col)
            formula = f'=SUMIFS(Registos!$G:$G,Registos!$A:$A,"{nc}",Registos!$D:$D,$A{linha})'
            c = ws.cell(row=linha, column=col, value=formula)
            c.number_format = FORMATO_MIN
        _linha_zebra(ws, linha, 1, 1 + len(nomes_cenarios))
        linha += 1
    linha_total = linha
    ws.cell(row=linha_total, column=1, value="Total").font = Font(bold=True, color=COR_TEXTO)
    for i, nc in enumerate(nomes_cenarios):
        col = 2 + i
        col_letra = get_column_letter(col)
        primeira = linha_cab1 + 1
        ultima = linha_total - 1
        c = ws.cell(row=linha_total, column=col, value=f"=SUM({col_letra}{primeira}:{col_letra}{ultima})")
        c.number_format = FORMATO_MIN
        c.font = Font(bold=True, color=COR_TEXTO)
    linha = linha_total + 2

    # (2) ritmo: máximo seguido, maior pausa, n.º de entradas por cenário.
    ws.cell(row=linha, column=1, value="Ritmo por atleta").font = Font(bold=True, color=COR_TEXTO)
    linha += 1
    titulos = ["Atleta"]
    for nc in nomes_cenarios:
        titulos += [f"{nc} — máx. seguido", f"{nc} — maior pausa", f"{nc} — n.º entradas"]
    _cabecalho(ws, linha, 1, titulos)
    linha += 1
    for nome in todos_atletas:
        ws.cell(row=linha, column=1, value=nome)
        col = 2
        for nc in nomes_cenarios:
            cfg, plano = cenarios[nc]
            minutos_jogo_total = cfg["jogo"]["partes"] * cfg["jogo"]["minutos_por_parte"]
            if nome in {a["nome"] for a in cfg["atletas"]}:
                stats = estatisticas_atleta(plano, nome, cfg["jogo"]["minutos_por_parte"], minutos_jogo_total)
                ws.cell(row=linha, column=col, value=stats["maximo_seguido"])
                ws.cell(row=linha, column=col + 1, value=stats["maior_pausa"] if stats["maior_pausa"] is not None else "—")
                ws.cell(row=linha, column=col + 2, value=stats["n_entradas"])
            col += 3
        _linha_zebra(ws, linha, 1, col - 1)
        linha += 1
    linha += 2

    # (3) minutos por posição, por cenário.
    for nc in nomes_cenarios:
        cfg, _ = cenarios[nc]
        ordem_vagas = cfg["posicoes"]
        ws.cell(row=linha, column=1, value=f"{nc} — minutos por posição").font = Font(bold=True, color=COR_TEXTO)
        linha += 1
        _cabecalho(ws, linha, 1, ["Atleta"] + ordem_vagas + ["Total"])
        linha += 1
        nomes_cenario = [a["nome"] for a in cfg["atletas"]]
        for nome in nomes_cenario:
            ws.cell(row=linha, column=1, value=nome)
            for i, vaga in enumerate(ordem_vagas):
                col = 2 + i
                col_letra = get_column_letter(col)
                formula = f'=SUMIFS(Registos!$G:$G,Registos!$A:$A,"{nc}",Registos!$D:$D,$A{linha},Registos!$C:$C,"{vaga}")'
                c = ws.cell(row=linha, column=col, value=formula)
                c.number_format = FORMATO_MIN
            col_total = 2 + len(ordem_vagas)
            primeira_letra = get_column_letter(2)
            ultima_letra = get_column_letter(col_total - 1)
            c = ws.cell(row=linha, column=col_total, value=f"=SUM({primeira_letra}{linha}:{ultima_letra}{linha})")
            c.number_format = FORMATO_MIN
            _linha_zebra(ws, linha, 1, col_total)
            linha += 1
        linha += 2

    for col, largura in zip("ABCDEFGHIJ", [16] + [15] * 9):
        ws.column_dimensions[col].width = largura


# --------------------------------------------------------------------------- Atletas


def escrever_atletas(ws: Worksheet, cenarios: dict[str, tuple[dict, dict]]) -> None:
    ws.title = "Atletas"
    _ajustar_pagina(ws)
    ws.cell(row=1, column=1, value="Atletas").font = Font(bold=True, size=14, color=COR_TEXTO)

    nomes_cenarios = list(cenarios.keys())
    todos_atletas: dict[str, dict] = {}
    for nc in nomes_cenarios:
        cfg, _ = cenarios[nc]
        for a in cfg["atletas"]:
            todos_atletas.setdefault(a["nome"], {})[nc] = a

    linha = 3
    _cabecalho(ws, linha, 1, ["Atleta", "Papel"] + [f"Posições — {nc}" for nc in nomes_cenarios])
    linha += 1
    for nome, por_cenario in todos_atletas.items():
        ws.cell(row=linha, column=1, value=nome)
        papel = next(iter(por_cenario.values()))["papel"]
        ws.cell(row=linha, column=2, value=papel)
        for i, nc in enumerate(nomes_cenarios):
            a = por_cenario.get(nc)
            ws.cell(row=linha, column=3 + i, value=", ".join(a["posicoes"]) if a else "—")
        _linha_zebra(ws, linha, 1, 2 + len(nomes_cenarios))
        linha += 1
    linha += 2

    ws.cell(row=linha, column=1, value="Parâmetros do jogo").font = Font(bold=True, color=COR_TEXTO)
    linha += 1
    _cabecalho(ws, linha, 1, ["Cenário", "Partes", "Minutos por parte", "Minutos de jogo"])
    linha_cab_param = linha
    linha += 1
    for nc in nomes_cenarios:
        cfg, _ = cenarios[nc]
        ws.cell(row=linha, column=1, value=nc)
        ws.cell(row=linha, column=2, value=cfg["jogo"]["partes"])
        ws.cell(row=linha, column=3, value=cfg["jogo"]["minutos_por_parte"])
        ws.cell(row=linha, column=4, value=f"=B{linha}*C{linha}")
        _linha_zebra(ws, linha, 1, 4)
        linha += 1
    linha += 2

    ws.cell(row=linha, column=1, value="Notas — regras usadas").font = Font(bold=True, color=COR_TEXTO)
    linha += 1
    for nc in nomes_cenarios:
        cfg, _ = cenarios[nc]
        ws.cell(row=linha, column=1, value=f"{nc}: {_texto_legenda(cfg)}")
        linha += 1

    for col, largura in zip("ABCDEFGH", [16, 12] + [22] * 6):
        ws.column_dimensions[col].width = largura


# --------------------------------------------------------------------------- orquestração


def exportar(cenarios: dict[str, tuple[dict, dict]], descricoes: dict[str, str], caminho_saida: Path) -> None:
    wb = Workbook()
    ws_resumo = wb.active
    escrever_resumo(ws_resumo, cenarios)

    ws_registos = wb.create_sheet("Registos")
    linha_de = escrever_registos(ws_registos, cenarios)

    for nc, (cfg, plano) in cenarios.items():
        ws = wb.create_sheet(nc)
        escrever_cenario(ws, nc, descricoes.get(nc, ""), cfg, plano, linha_de)

    escrever_atletas(wb.create_sheet("Atletas"), cenarios)

    ordem = ["Resumo"] + list(cenarios.keys()) + ["Atletas", "Registos"]
    wb._sheets.sort(key=lambda ws: ordem.index(ws.title))

    caminho_saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(caminho_saida)


def recalcular_com_libreoffice(caminho: Path) -> None:
    """Recalcula fórmulas gravando os valores (secção 7/8 da SPEC.md) — exige
    LibreOffice instalado (`soffice`); se não encontrar, avisa e segue sem recalcular.

    `--convert-to` pro MESMO diretório do ficheiro de origem falha (LibreOffice tenta
    gravar por cima de um ficheiro que ainda está "em uso" pela própria conversão) —
    converte pra um diretório temporário e move o resultado de volta por cima do
    original, só depois da conversão terminar com sucesso.
    """
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, str(caminho)],
                check=True,
                capture_output=True,
                timeout=120,
            )
        except FileNotFoundError:
            print("⚠️  LibreOffice (soffice) não encontrado — fórmulas não foram recalculadas. Abra o Excel manualmente pra ver os valores.", file=sys.stderr)
            return
        except subprocess.CalledProcessError as e:
            print(f"⚠️  LibreOffice falhou ao recalcular: {e.stderr.decode(errors='replace')}", file=sys.stderr)
            return
        recalculado = Path(tmp) / caminho.name
        if recalculado.exists():
            shutil.move(str(recalculado), str(caminho))
        else:
            print(f"⚠️  LibreOffice não gerou {recalculado.name} — fórmulas não foram recalculadas.", file=sys.stderr)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("saida", help="Caminho do .xlsx de saída")
    ap.add_argument("pares", nargs="+", help="config.json:plano_resolvido.json, um por cenário")
    args = ap.parse_args()

    cenarios: dict[str, tuple[dict, dict]] = {}
    descricoes: dict[str, str] = {}
    for par in args.pares:
        caminho_cfg, caminho_plano = par.split(":")
        cfg = json.loads(Path(caminho_cfg).read_text(encoding="utf-8"))
        plano = json.loads(Path(caminho_plano).read_text(encoding="utf-8"))
        nome = Path(caminho_cfg).stem
        cenarios[nome] = (cfg, plano)
        descricoes[nome] = cfg.get("_descricao", nome)

    caminho_saida = Path(args.saida)
    exportar(cenarios, descricoes, caminho_saida)
    recalcular_com_libreoffice(caminho_saida)
    print(f"✅ Excel gravado em {caminho_saida}")


if __name__ == "__main__":
    main()
